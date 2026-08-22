from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from reporting.models import SecurityReport

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
_BODY_FONT = Font(name="Arial")
_SEVERITY_FILLS = {
    "error": PatternFill("solid", fgColor="F4CCCC"),
    "warning": PatternFill("solid", fgColor="FCE5CD"),
    "info": PatternFill("solid", fgColor="D9EAF7"),
    "unknown": PatternFill("solid", fgColor="E7E6E6"),
}


def _style_header(sheet: Worksheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_body(sheet: Worksheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if not cell.font.bold:
                cell.font = _BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _configure_print(sheet: Worksheet, *, landscape: bool = False) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5


def render_excel(report: SecurityReport) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Security Report"])
    summary.append(["Generated At", report.generated_at])
    summary.append(["Total Findings", len(report.findings)])
    summary.append([])
    summary.append(["Severity", "Count"])
    for severity in ("error", "warning", "info", "unknown"):
        summary.append([severity, report.counts[severity]])
    summary.merge_cells("A1:B1")
    summary["A1"].font = Font(name="Arial", bold=True, size=16, color="FFFFFF")
    summary["A1"].fill = _HEADER_FILL
    summary.row_dimensions[1].height = 28
    _style_header(summary, 5)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 32

    sources = workbook.create_sheet("Sources")
    sources.append(["Tool", "Version", "Ruleset", "Target", "Started At"])
    for source in report.sources:
        sources.append(
            [source.tool, source.tool_version, source.ruleset, source.target, source.started_at]
        )
    _style_header(sources)
    sources.freeze_panes = "A2"
    sources.auto_filter.ref = sources.dimensions
    for column, width in zip("ABCDE", (18, 16, 28, 40, 26), strict=True):
        sources.column_dimensions[column].width = width

    findings = workbook.create_sheet("Findings")
    findings.append(
        [
            "ID",
            "Severity",
            "Tool",
            "Rule ID",
            "Path",
            "Start Line",
            "Message",
            "Code",
            "Metadata",
            "Raw Reference",
        ]
    )
    for item in report.findings:
        findings.append(
            [
                item["id"],
                item["severity"],
                item["tool"],
                item["rule_id"],
                item["path"],
                item["start_line"],
                item["message"],
                item.get("code"),
                json.dumps(item.get("metadata", {}), ensure_ascii=False, sort_keys=True),
                item["raw_reference"],
            ]
        )
        findings.cell(findings.max_row, 2).fill = _SEVERITY_FILLS[item["severity"]]
    _style_header(findings)
    findings.freeze_panes = "A2"
    findings.auto_filter.ref = findings.dimensions
    widths = (20, 12, 16, 28, 40, 12, 54, 54, 48, 24)
    for column, width in zip("ABCDEFGHIJ", widths, strict=True):
        findings.column_dimensions[column].width = width

    for sheet in workbook.worksheets:
        _style_body(sheet)
    _configure_print(summary)
    _configure_print(sources, landscape=True)
    _configure_print(findings, landscape=True)
    sources.print_title_rows = "1:1"
    findings.print_title_rows = "1:1"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
