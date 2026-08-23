import unittest
from io import BytesIO

from openpyxl import load_workbook

from agent.models import AnalysisDocument, AnalysisItem
from reporting.models import ScanSource, SecurityReport
from reporting.render_excel import render_excel


class ExcelRendererTests(unittest.TestCase):
    def test_excel_contains_summary_sources_and_findings(self) -> None:
        source = ScanSource("semgrep", "1", "rules", ".", "2026-08-22T00:00:00Z")
        finding = {
            "id": "id",
            "tool": "semgrep",
            "rule_id": "rule",
            "severity": "error",
            "path": "a.py",
            "start_line": 2,
            "message": "Avoid exec.",
            "code": "exec(x)",
            "metadata": {"cwe": ["CWE-78"]},
            "raw_reference": "/results/0",
        }
        report = SecurityReport.create([source], [finding])

        workbook = load_workbook(BytesIO(render_excel(report)), data_only=False)

        self.assertEqual(workbook.sheetnames, ["Summary", "Sources", "Findings"])
        self.assertEqual(workbook["Summary"]["B3"].value, 1)
        self.assertTrue(workbook["Summary"]["A1"].font.bold)
        self.assertGreaterEqual(workbook["Summary"].row_dimensions[1].height or 0, 24)
        self.assertEqual(workbook["Sources"]["A2"].value, "semgrep")
        self.assertEqual(workbook["Findings"]["A2"].value, "id")
        self.assertEqual(workbook["Findings"]["G2"].value, "Avoid exec.")
        self.assertEqual(workbook["Sources"].page_setup.fitToWidth, 1)
        self.assertEqual(workbook["Findings"].page_setup.orientation, "landscape")

    def test_excel_renders_analysis_columns(self) -> None:
        source = ScanSource("semgrep", "1", "rules", ".", "2026-08-22T00:00:00Z")
        finding = {
            "id": "id",
            "tool": "semgrep",
            "rule_id": "rule",
            "severity": "error",
            "path": "a.py",
            "start_line": 2,
            "message": "Avoid exec.",
            "code": "exec(x)",
            "metadata": {"cwe": ["CWE-78"]},
            "raw_reference": "/results/0",
        }
        report = SecurityReport.create([source], [finding])
        analysis = AnalysisDocument.create(
            "deterministic",
            [
                AnalysisItem(
                    "id", "confirmed", "OS command injection", "cause",
                    "impact", "remediation", ("CWE-78",), diff_status="changed",
                )
            ],
        )

        workbook = load_workbook(BytesIO(render_excel(report, analysis)), data_only=False)
        sheet = workbook["Findings"]
        self.assertEqual(sheet["K1"].value, "Classification")
        self.assertEqual(sheet["L1"].value, "Diff Status")
        self.assertEqual(sheet["M1"].value, "Baseline Status")
        self.assertEqual(sheet["K2"].value, "confirmed")
        self.assertEqual(sheet["L2"].value, "changed")
        self.assertEqual(sheet["N2"].value, "OS command injection")
        self.assertEqual(sheet["O2"].value, "cause")


if __name__ == "__main__":
    unittest.main()
